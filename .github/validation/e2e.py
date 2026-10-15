"""Synthetic, independent pre-PR smoke checks; no personal input files."""
import json, os, subprocess, sys, zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageChops
from pypdf import PdfReader
from pypdf.generic import IndirectObject

cli = str(Path(sys.argv[1]).resolve())
root = Path("target/validation-output").resolve()
root.mkdir(exist_ok=True)
def run(*args):
    p = subprocess.run([cli, *map(str,args)], text=True, capture_output=True, timeout=180)
    if p.returncode:
        raise RuntimeError(f"{args}: {p.stdout}\n{p.stderr}")
    return p.stdout
def steps(name, seq):
    path = root / (name + ".json")
    path.write_text(json.dumps(seq), encoding="utf8")
    output = run("run", "--script", path, "--root", root)
    (root / (name + "-results.json")).write_text(output, encoding="utf8")
def call(tool, **args):
    return {"tool":tool,"args":args}
steps("text", [
    call("doc_create", **{"from":"text", "text":"TESTE DE OCULTACAO\nSEGREDO123\nPUBLICO 456"}),
    call("doc_save",doc=1,path="original.pdf"),
    call("redact_mark",doc=1,find="SEGREDO123"),
    call("doc_save",doc=1,path="marked.pdf"),
    call("redact_apply",doc=1),
    call("doc_save",doc=1,path="applied.pdf"),
    call("doc_close",doc=1),
    call("doc_open",path="applied.pdf"),
    call("text_extract",doc=2),
])
assert "SEGREDO123" in PdfReader(root/"marked.pdf").pages[0].extract_text()
txt = PdfReader(root/"applied.pdf").pages[0].extract_text()
assert "SEGREDO123" not in txt and "PUBLICO 456" in txt, txt
for ext in ["jpg","png"]:
    im=Image.new("RGB",(600,400),"white")
    d=ImageDraw.Draw(im)
    d.rectangle((40,40,200,200),fill=(210,30,30))
    d.rectangle((350,50,550,300),fill=(20,60,220))
    im.save(root/("scan."+ext))
    steps(ext, [
        call("doc_create", **{"from":"images","paths":["scan."+ext],"dpi":72}),
        call("doc_save",doc=1,path=ext+"-before.pdf"),
        call("redact_mark",doc=1,page=1,rect=[60,60,120,120]),
        call("redact_apply",doc=1),
        call("doc_save",doc=1,path=ext+"-after.pdf"),
        call("doc_close",doc=1),
        call("doc_open",path=ext+"-after.pdf"),
        {"tool":"page_render","args":{"doc":2,"page":1,"dpi":72},"out":ext+"-after.png"},
    ])
    before=PdfReader(root/(ext+"-before.pdf"))
    after=PdfReader(root/(ext+"-after.pdf"))
    original=list(before.pages[0].images)[0].image.convert("RGB")
    images=list(after.pages[0].images)
    assert len(images)==1, (ext,len(images))
    changed=images[0].image.convert("RGB")
    assert changed.size == original.size
    assert changed.crop((65,65,115,115)).getbbox() is None, "covered samples remain"
    # Independent JPEG decoders use different chroma upsampling at sharp edges.
    # Compare the solid interior; Rust regression tests assert every uncovered sample exactly.
    delta=ImageChops.difference(original.crop((370,70,530,280)),changed.crop((370,70,530,280)))
    assert max(hi for lo,hi in delta.getextrema()) <= 4, "public blue block changed"
    old_stream=next(iter(before.pages[0]["/Resources"]["/XObject"].values())).get_object()
    original_samples=old_stream.get_data()
    for generation, entries in after.xref.items():
        for number in entries:
            if number == 0: continue
            obj=after.get_object(IndirectObject(number,generation,after))
            if hasattr(obj,"get_data") and obj.get("/Subtype")=="/Image":
                assert obj.get_data()!=original_samples, "original image remains recoverable"
    assert Image.open(root/(ext+"-after.png")).convert("RGB").getpixel((400,100))[2] > 150
steps("combine",[
    call("doc_combine",paths=["applied.pdf","jpg-after.pdf","png-after.pdf"]),
    call("doc_save",doc=1,path="combined.pdf"),
    call("doc_export_office",doc=1,path="combined.docx"),
    call("doc_export_office",doc=1,path="combined.pptx"),
])
assert len(PdfReader(root/"combined.pdf").pages)==3
for name, member in [("combined.docx","word/document.xml"),("combined.pptx","ppt/presentation.xml")]:
    with zipfile.ZipFile(root/name) as z:
        assert z.testzip() is None and member in z.namelist()
steps("images",[
    call("doc_create",**{"from":"images","paths":["scan.jpg","scan.png"],"dpi":72}),
    call("doc_save",doc=1,path="images.pdf"),
])
assert len(PdfReader(root/"images.pdf").pages)==2
run("check",root/"combined.pdf",root/"images.pdf","--timeout","30","--json",root/"robustness.json")
assert all(item["status"] == "ok" for item in json.loads((root/"robustness.json").read_text()))
(root/"PASS.txt").write_text("PASS: text mark/apply/save/reopen; JPEG and PNG redaction, retired original image objects; rendered public pixels; combine PDFs; combine images; DOCX/PPTX export ZIP integrity.\n",encoding="utf8")
print((root/"PASS.txt").read_text())

# Genuine Office import fixtures, created here from original text, shapes and pixels.
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches as PptInches
from openpyxl import Workbook
from openpyxl.styles import Font
doc=Document()
doc.styles["Normal"].font.name="Inter"
doc.add_heading("PUBLICO - Word conversion",0)
doc.add_paragraph("Acentos: ação, coração, São Paulo. Two-page synthetic conversion.")
table=doc.add_table(rows=2,cols=2)
table.cell(0,0).text="Item"
table.cell(0,1).text="Valor"
table.cell(1,0).text="Exemplo"
table.cell(1,1).text="42"
doc.add_picture(str(root/"scan.png"),width=DocInches(4))
doc.add_page_break()
doc.add_paragraph("PUBLICO - second page")
doc.save(root/"office.docx")
prs=Presentation()
for number in [1,2]:
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    box=slide.shapes.add_textbox(PptInches(1),PptInches(0.5),PptInches(8),PptInches(1))
    box.text_frame.text=f"PUBLICO - Slide {number}"
    for paragraph in box.text_frame.paragraphs:
        for run_ in paragraph.runs: run_.font.name="Inter"
    slide.shapes.add_picture(str(root/"scan.png"),PptInches(1),PptInches(2),width=PptInches(6))
prs.save(root/"office.pptx")
wb=Workbook()
ws=wb.active
ws.append(["PUBLICO", "Valor"])
ws.append(["Exemplo",42])
for row in ws:
    for cell in row: cell.font=Font(name="Inter")
ws.column_dimensions["A"].width=24
wb.save(root/"office.xlsx")
lo=os.environ.get("PDFCRAFT_LIBREOFFICE") or (r"C:\Program Files\LibreOffice\program\soffice.com" if os.name=="nt" else "/usr/bin/libreoffice")
for src, exts in [("docx",["doc","odt","rtf"]),("pptx",["ppt","odp"]),("xlsx",["xls","ods"])]:
    for ext in exts:
        subprocess.run([lo, "-env:UserInstallation="+(root/("fixture-profile-"+ext)).as_uri(), "--headless","--convert-to",ext,"--outdir",str(root),str(root/("office."+src))],check=True,timeout=120,capture_output=True)
        assert (root/("office."+ext)).exists(), ext
for ext in ["docx","doc","pptx","ppt","xlsx","xls","odt","odp","ods","rtf"]:
    source=root/("office."+ext)
    original_bytes=source.read_bytes()
    script_path=root/("import-"+ext+".json")
    script_path.write_text(json.dumps([
        call("doc_convert_office",path=str(source)),
        call("doc_save",doc=1,path=str(root/("converted-"+ext+".pdf"))),
        {"tool":"page_render","args":{"doc":1,"page":1,"dpi":72},"out":str(root/("converted-"+ext+".png"))},
    ]),encoding="utf8")
    result=run("run","--script",script_path)
    (root/("import-"+ext+"-results.json")).write_text(result,encoding="utf8")
    converted=PdfReader(root/("converted-"+ext+".pdf"))
    assert "PUBLICO" in "".join(page.extract_text() for page in converted.pages),ext
    if ext in ["docx","doc","odt","rtf","pptx","ppt","odp"]:
        assert len(converted.pages)==2,(ext,len(converted.pages))
    assert source.read_bytes()==original_bytes
(root/"OFFICE-PASS.txt").write_text("PASS: DOCX/DOC/PPTX/PPT/XLSX/XLS/ODT/ODP/ODS/RTF to PDF through the app CLI; source unchanged; PDF reopened independently; public text and page counts checked.\n",encoding="utf8")
print((root/"OFFICE-PASS.txt").read_text())
